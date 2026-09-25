"""Gate-status report (derived, receipted, verifiable) and the workpaper that renders it (quality policy §2, D16).

Every test redirects the policy, approvals and run records to temporary copies.
"""
import json
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from tests.test_recurring import ROOT, _cli, _load, _attest
from tests.test_governance_events import _keys, _approve, _authorise, _review, _signed_mapping_with_demo_keys

SGT = timezone(timedelta(hours=8))


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    from governance.production import policy_approval as pa, gate_status as gs
    copy = tmp_path / "policy/quality_policy.md"
    copy.parent.mkdir()
    shutil.copyfile(pa.POLICY, copy)
    monkeypatch.setattr(pa, "POLICY", copy)
    monkeypatch.setattr(pa, "APPROVALS", tmp_path / "policy/approvals")
    register = tmp_path / "policy/unexercised_guards.md"
    shutil.copyfile(gs.REGISTER, register)
    monkeypatch.setattr(gs, "REGISTER", register)
    monkeypatch.setattr(gs, "RUNS", tmp_path / "runs")
    monkeypatch.setattr(gs, "PACKS", tmp_path / "packs")
    return tmp_path


def _governed_series(tmp_path):
    home = tmp_path / "home"
    home.mkdir()
    keys = _keys(home)
    approval = _approve(keys)
    document = _signed_mapping_with_demo_keys(home, tmp_path)
    review = _review(keys, document, "outsider.key", "Independent Reviewer")
    result = _authorise(home, policy_approval=approval["approval"], mapping=str(document),
                        mapping_review=review["review"])
    config_path = Path(result["config"])
    config, root = _load(config_path)
    from governance.production import recurring
    _cli(home, "demo-inbox", "--config", config_path, "--week", 3)
    recurring.tick(config, root, now=datetime(2026, 10, 1, 9, tzinfo=SGT))
    _attest(config, root, "CHG-WEEKLY-2026-09-29", "NO_EXCEPTIONS_NOTED", True,
            "Week 3 checked with corroborated coverage; nothing raised.")
    return config_path, approval


def _states(report):
    return {g["gate"]: g["state"] for g in report["gates"]}


def test_the_report_is_derived_from_events_and_carries_its_receipt(isolated):
    from governance.production import gate_status as gs
    config_path, approval = _governed_series(isolated)
    report = gs.generate(str(config_path))
    states = _states(report)
    assert states["Policy approval"] == "HELD" and states["Independent mapping review"] == "HELD"
    assert states["Attestation quality"] == "HELD" and states["Procedure completeness"] == "OPEN"
    assert states["A valid test run"] == "OPEN"                   # no run record in this isolated folder
    kinds = {i["kind"] for i in report["derivation_receipt"]["inputs"]}
    assert {"governing policy", "policy approval", "mapping review", "case journal", "guard register"} <= kinds
    assert gs.verify(report)["valid"] is True


def test_a_hand_edited_report_is_refused_even_with_its_hash_recomputed(isolated):
    from governance.production import gate_status as gs
    from governance.investigation.store import digest
    report = gs.generate()
    edited = json.loads(json.dumps(report))
    edited["gates"][1]["state"] = "HELD"                          # procedure completeness, claimed
    assert "it was edited" in gs.verify(edited)["reason"]
    body = {k: edited[k] for k in ("report", "format", "policy", "series", "as_of", "live", "gates", "guard_register")}
    edited["body_sha256"] = digest(body)                           # a careful forger recomputes the hash
    assert "gives a different result" in gs.verify(edited)["reason"]
    assert gs.verify({k: v for k, v in report.items() if k != "derivation_receipt"})["reason"].startswith(
        "no derivation receipt")


def test_a_changed_input_is_named_and_later_events_do_not_invalidate_a_report(isolated):
    from governance.production import gate_status as gs
    report = gs.generate()
    keys = _keys(isolated)
    _approve(keys)                                                 # happens after the report's moment
    assert gs.verify(report)["valid"] is True, "later events must not change a report about an earlier moment"
    gs.REGISTER.write_text(gs.REGISTER.read_text() + "\n")         # an input the report did read
    assert gs.verify(report)["reason"] == "an input named in the receipt has changed"


def test_historical_status_is_derived_with_as_of(isolated):
    from governance.production import gate_status as gs
    approval = _approve(_keys(isolated))
    approved_at = gs._when(json.loads(Path(approval["approval"]).read_text())["payload"]["approved_at"])
    before = gs.generate(as_of=(approved_at - timedelta(minutes=1)).isoformat())
    after = gs.generate(as_of=(approved_at + timedelta(minutes=1)).isoformat())
    assert _states(before)["Policy approval"] == "OPEN" and _states(after)["Policy approval"] == "HELD"
    assert before["live"] is False and gs.verify(before)["valid"] is True


def test_run_and_trace_records_feed_their_gates(isolated):
    from governance.production import gate_status as gs
    gs.RUNS.mkdir()
    now = datetime.now().astimezone()
    (gs.RUNS / "run.json").write_text(json.dumps({"at": now.isoformat(), "verdict": "COMPLETE RUN, ALL PASSED",
                                                  "collected": 10, "passed": 10, "skipped": 0,
                                                  "environment_missing": []}))
    (gs.RUNS / "trace-x.json").write_text(json.dumps({"at": now.isoformat(), "verdict": "PASS", "unreached": 12,
                                                      "registered": 12}))
    states = _states(gs.generate())
    assert states["A valid test run"] == "HELD"
    assert states["A refusal or integrity rule"] == "PARTIAL"      # passing, but 12 guards still unreached


def test_the_policy_carries_rules_and_points_to_the_report(isolated):
    from governance.production import policy_approval as pa
    text = pa.POLICY.read_text()
    assert "How it is held today" not in text and "Not yet executed" not in text
    assert "derivation receipt" in text and "tools/gaar_gate_status.py verify" in text
    assert "Sampled blind readings" in text and "Evidence first" in text
    assert "D16" in text


def test_the_workpaper_renders_only_a_verified_report_about_this_exact_policy(isolated):
    import importlib
    from docx import Document
    from governance.production import gate_status as gs, policy_approval as pa
    tool = importlib.import_module("tools.gaar_workpaper")
    config_path, approval = _governed_series(isolated)
    report = gs.generate(str(config_path))
    path = isolated / "report.json"
    path.write_text(json.dumps(report))
    out = tool.build(path, isolated / "workpaper.docx")
    doc = Document(out["workpaper"])
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert pa.policy_identity(pa.POLICY)["policy_sha256"] in text     # the full hash in the banner
    assert "Appendix A — Gate status" in text and "Derivation receipt" in text
    assert "HELD: policy 1.3" in text or "HELD" in text
    section9 = text[text.index("9. Vocabulary"):text.index("10. Independence")]
    assert "1.\u2002Changes to decision" in section9                   # the policy's own numbering

    edited = json.loads(path.read_text())
    gate = next(g for g in edited["gates"] if g["state"] == "OPEN")     # a real change, not a no-op
    gate["state"] = "HELD"
    path.write_text(json.dumps(edited))
    with pytest.raises(ValueError, match="does not verify"):
        tool.build(path, isolated / "refused.docx")
    path.write_text(json.dumps(report))
    pa.POLICY.write_text(pa.POLICY.read_text() + "\nA later edit.\n")
    with pytest.raises(ValueError):
        tool.build(path, isolated / "refused2.docx")


def test_the_refusal_gate_is_held_only_when_no_guard_is_unreached(isolated):
    from governance.production import gate_status as gs
    gs.RUNS.mkdir()
    now = datetime.now().astimezone()
    (gs.RUNS / "trace-x.json").write_text(json.dumps({"at": now.isoformat(), "verdict": "PASS", "unreached": 0,
                                                      "registered": 0}))
    assert _states(gs.generate())["A refusal or integrity rule"] == "HELD"
    (gs.RUNS / "trace-y.json").write_text(json.dumps({"at": (now + timedelta(seconds=1)).isoformat(),
                                                      "verdict": "FAIL", "unreached": 1, "registered": 0}))
    assert _states(gs.generate(as_of=(now + timedelta(seconds=2)).isoformat()))["A refusal or integrity rule"] == "OPEN"


def test_a_series_assessed_on_a_simulated_clock_says_so(isolated):
    """Seen on the Mac: live report 2026-09-24 read '3 period(s) assessed, 2 due' with no reason given."""
    from governance.production import gate_status as gs
    config_path, _ = _governed_series(isolated)
    evidence = next(g["evidence"] for g in gs.generate(str(config_path))["gates"] if g["gate"] == "Collection completeness")
    assert "simulated clock (latest 2026-10-01T09:00:00+08:00)" in evidence


def test_the_workpaper_keeps_each_wrapped_bullet_whole_and_ends_without_a_blank_page(isolated):
    """Seen on the Mac rendering of v18: a bullet wrapped in the source became a bullet plus a stray paragraph
    (§2, §5, §11), and a 10th blank page followed the receipt table."""
    import importlib
    import re
    from docx import Document
    from docx.oxml.ns import qn
    from governance.production import gate_status as gs, policy_approval as pa
    tool = importlib.import_module("tools.gaar_workpaper")
    report = gs.generate()
    path = isolated / "report.json"
    path.write_text(json.dumps(report))
    doc = Document(tool.build(path, isolated / "workpaper.docx")["workpaper"])
    paragraphs = [p.text for p in doc.paragraphs]
    source = pa.POLICY.read_text().splitlines()
    bullets = 0
    for n, line in enumerate(source):
        if line.startswith("- "):
            item = [line[2:]]
            while n + 1 < len(source) and source[n + 1].startswith("  ") and source[n + 1].strip():
                n += 1
                item.append(source[n].strip())
            expected = re.sub(r"\*\*|`", "", " ".join(item))
            assert expected in paragraphs, expected
            bullets += 1
    assert bullets >= 10
    assert not any(p.startswith("the canonical test and trace records") for p in paragraphs)
    last = doc.paragraphs[-1]
    assert last.text == "" and last._p.pPr.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")) == "2"


def test_a_run_that_stopped_for_its_environment_keeps_the_gate_open_and_the_report_readable(isolated):
    # D27 (v28 round): the record written by ENVIRONMENT NOT READY has no pass counts; the report crashed reading it,
    # so the scheduler's gate_status job failed on every tick until a full run was recorded.
    from governance.production import gate_status as gs
    gs.RUNS.mkdir()
    now = datetime.now().astimezone()
    (gs.RUNS / "run.json").write_text(json.dumps({"at": now.isoformat(), "verdict": "ENVIRONMENT NOT READY",
                                                  "collected": 0, "executed": 0, "conda_env": "base",
                                                  "environment_missing": ["streamlit==1.63.0"]}))
    report = gs.generate()
    gate = next(g for g in report["gates"] if g["gate"] == "A valid test run")
    assert gate["state"] == "OPEN"
    assert "stopped before any test: environment not ready (1 package(s) missing in base)" in gate["evidence"]
